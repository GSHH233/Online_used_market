# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"F:\Online_used_marketplace")
INPUT = ROOT / "resume_input.docx"
OUTPUT = Path(r"F:\大三下\面试\个人简历_AI应用开发优化版_v2.docx")


def set_run_font(run, size: float, bold: bool = False, color: RGBColor | None = None) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if color is not None:
        run.font.color.rgb = color


def remove_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    set_run_font(run, 10.2, bold=True, color=RGBColor(31, 78, 121))

    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "9E9E9E")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0.4)
    p.paragraph_format.line_spacing = 0.88
    run = p.add_run("• " + text)
    set_run_font(run, 8.35)


def add_line(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0.3)
    p.paragraph_format.line_spacing = 0.9
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_run_font(run, 8.35, bold=True)
        run = p.add_run(text[len(bold_prefix):])
        set_run_font(run, 8.35)
    else:
        run = p.add_run(text)
        set_run_font(run, 8.35)


def add_project_title(doc: Document, left: str, right: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(0.3)
    p.paragraph_format.line_spacing = 0.9
    run = p.add_run(left)
    set_run_font(run, 9.1, bold=True)
    run = p.add_run("\t" + right)
    set_run_font(run, 9.1, bold=True)


def read_header_rows() -> list[list[str]]:
    original = Document(INPUT)
    if not original.tables:
        return [
            ["性别：男", "年龄：20", ""],
            ["电话：", "邮箱：", "专业：软件工程"],
            ["意向岗位：AI应用开发 / Agent开发实习", "学校：", "毕业时间：2027.6"],
        ]
    rows: list[list[str]] = []
    for row in original.tables[0].rows:
        rows.append([cell.text.replace("\n", " ").strip() for cell in row.cells])
    if len(rows) >= 3 and rows[2]:
        rows[2][0] = "意向岗位：AI应用开发 / Agent开发实习"
    return rows


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.42)
    section.right_margin = Inches(0.42)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(8.4)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(2)
    run = name.add_run("1")
    set_run_font(run, 15, bold=True)

    header_rows = read_header_rows()
    max_cols = max(len(row) for row in header_rows)
    table = doc.add_table(rows=len(header_rows), cols=max_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_idx, row in enumerate(header_rows):
        for col_idx in range(max_cols):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            remove_cell_border(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(row[col_idx] if col_idx < len(row) else "")
            set_run_font(run, 8.2)

    add_heading(doc, "专业技能")
    for item in [
        "AI应用：熟悉 Spring AI、Prompt 编排、RAG、ReAct/Graph Agent、Tool Use、流式输出（WebSocket/SSE）与多轮上下文管理。",
        "RAG工程：掌握文档切片、Embedding、Milvus 向量检索、Elasticsearch BM25、Hybrid + RRF 融合、metadata 过滤与检索评测。",
        "后端开发：熟悉 Java / Spring Boot / MyBatis-Plus 分层开发，具备 RESTful API、任务状态表、统一结果模型与异常处理实践。",
        "中间件与数据：熟悉 MySQL、Redis、RabbitMQ、Docker、Prometheus，能进行缓存、消息、日志、指标与服务联调。",
        "工程习惯：熟悉 Git、Linux 基础操作、Postman 接口调试，具备评测脚本、故障注入、重试降级与可观测性建设经验。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "项目经历")
    add_project_title(doc, "2025.11 - 至今    回血集市（二手交易 AI 助手）", "核心后端 / AI应用开发")
    add_line(
        doc,
        "技术栈：Spring Boot、Spring AI、MySQL、Redis、RabbitMQ、MyBatis-Plus、WebSocket、Milvus、Elasticsearch、Prometheus、Docker",
        "技术栈：",
    )
    for item in [
        "RAG知识库链路：实现文档入库、切片、Embedding、Milvus 向量检索与 ES 关键词检索，支持 vector-only / keyword-only / hybrid + RRF；处理 metadata/riskLevel 对齐、向量维度一致性、ES index 初始化和 reindex 幂等问题。",
        "Agent编排：实现 RAG Chat、ReAct Agent 与 Graph/DAG Agent，支持问题理解、工具选择、调用、Observation 汇总与最终回答；封装 knowledge_search、mysql_read、log_query、prometheus_query、vector_debug 等工具，并区分 USER / INTERNAL 模式。",
        "记忆系统：参考 mem0 思路实现 L1/L2/L3 分层记忆，接入 L2 语义召回、L3 KV Session Summary 与 Profile 偏好注入；设计 store、recall、consolidation、context assembly、graph expansion 五层记忆评测。",
        "可靠性与可观测：实现 Tool Execution Harness，支持超时控制、失败重试、Redis 快照恢复、fallback 降级、Micrometer/Prometheus 指标与故障注入评测，提升工具调用稳定性。",
        "工程闭环：沉淀 retrieval、rag_chat、agent、react、aiops、memory、harness、storage consistency 等评测脚本；新增 RAG Ingestion Job、Repair API 与任务状态表，记录阶段、DB/ES/Milvus 计数、一致性和耗时，支持批量修复。",
    ]:
        add_bullet(doc, item)

    add_project_title(doc, "2025.03 - 2025.06    凌霄优仓", "后端开发")
    add_line(doc, "技术栈：Spring Boot、Vue3、Nginx、Redis、MySQL、WebSocket、MyBatis-Plus", "技术栈：")
    for item in [
        "基于 WebSocket 实现用户与商家实时对话，设计消息格式、连接状态与断线重连逻辑，保障双端交互稳定。",
        "实现短信登录模块，使用 Redis 缓存验证码并控制有效期，处理重复提交、超时失效与登录安全校验。",
        "配置 Nginx 反向代理与负载均衡，结合 Redis 空值缓存与过期时间抖动，缓解缓存穿透和缓存雪崩风险。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "个人总结")
    add_bullet(doc, "具备 Java 后端与 AI 应用工程复合实践，能独立完成从接口、存储、工具编排到评测回归的功能闭环，适配中小团队快速迭代。")

    if OUTPUT.exists():
        OUTPUT.unlink()
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
